"""Equipo por mesa y jefe, fotos de Drive y desempeño de PRODUCTIVIDAD."""
import io
import json
import base64
import html
import hashlib
import hmac
import re
import unicodedata
from pathlib import Path
from urllib.parse import quote

import streamlit as st
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload
from PIL import Image, ImageOps

CARPETA_PREDETERMINADA = "1IZXOH9CeQk1XjNq3uvwWl2TkoKMNihcj"
PRODUCTIVIDAD_PREDETERMINADA = "1G4lAKm4fxV6ZsPoKyN49Kt5-Ub__hCm6YTvR4GBgIGw"
ESTRUCTURA_CARPETA_PREDETERMINADA = "1zQmGDrKYU4yl_Uyzy2bk9mOrnWbaAx80"
ACCESOS_PREDETERMINADO = "1RHAWRHEIeRrhJFYzZTzGBGg8GUcR9wxGgwYPrbA7Xs4"
LOGIN_HOJA_PREDETERMINADA = "LOGIN"
SCOPE = ["https://www.googleapis.com/auth/drive.readonly",
         "https://www.googleapis.com/auth/spreadsheets.readonly"]
FORMATOS = {"image/jpeg", "image/png", "image/webp", "image/gif"}
MAX_BYTES = 10 * 1024 * 1024
ORDEN_ORIGINAL = ["MESA", "ASESOR", "BOLSAS MEDIO DIA", "BOLSAS CIERRE",
                  "CLIENTES RECUPERADO", "CLIENTES NUEVO", "NO RENOVADOS",
                  "VARIACION CLIENTES", "VARIACIÓN DE CARTERA", "RECUPERACIÓN DE EFECTIVO",
                  "RECUPERACIÓN DE MORA (30%)", "RECUPERACION MORA AGOSTO", "MORA SEPTIEMBRE"]
NOMBRES_ASESOR = {"asesor", "asesores", "colaborador", "colaboradores", "personal",
                  "nombre", "nombre completo"}
SHEET_MIME = "application/vnd.google-apps.spreadsheet"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
INDICADORES_PRINCIPALES = [
    ("Bolsas al cierre", {"bolsas cierre"}),
    ("Clientes nuevos", {"clientes nuevo", "clientes nuevos"}),
    ("Clientes recuperados", {"clientes recuperado", "clientes recuperados"}),
]


def clave_mesa(mesa):
    clave = re.sub(r"^mesa\s*", "", normalizar(mesa))
    return str(int(float(clave))) if re.fullmatch(r"\d+(?:\.0+)?", clave) else clave


def interpretar_estructura(filas):
    """Soporta columnas de asesores/promotores y tablas con nombre/rol."""
    alias = {
        "mesa": {"mesa", "mesas"},
        "jefe": {"jefe de credito", "jefe de creditos", "jefe credito", "jefe"},
        "asesor": {"asesor", "asesores"},
        "promotor": {"promotor", "promotores"},
        "nombre": {"nombre", "nombre completo", "colaborador", "colaboradores", "personal", "asesor"},
        "rol": {"rol", "cargo", "puesto"},
    }
    indices, inicio = {}, None
    for i, fila in enumerate(filas[:10]):
        titulos = [normalizar(v) for v in fila]
        hallados = {k: next((j for j, titulo in enumerate(titulos) if titulo in nombres), None)
                    for k, nombres in alias.items()}
        largo = hallados["nombre"] is not None and hallados["rol"] is not None
        ancho = hallados["asesor"] is not None or hallados["promotor"] is not None
        if hallados["mesa"] is not None and hallados["jefe"] is not None and (largo or ancho):
            indices, inicio = hallados, i
            break
    if inicio is None:
        raise ValueError("ESTRUCTURA DE COLABORADORES necesita MESA/MESAS, JEFE DE CREDITO y ASESOR/ASESORES o PROMOTOR/PROMOTORES. También se admite NOMBRE y ROL.")
    def valor(fila, campo):
        j = indices.get(campo)
        return str(fila[j]).strip() if j is not None and j < len(fila) and fila[j] is not None else ""
    resultado, vistos, avisos = [], set(), []
    mesa_actual = jefe_actual = ""
    sin_mesa = sin_jefe = 0
    def agregar(nombre, rol, mesa, jefe):
        for persona in re.split(r"[\n;]+", nombre):
            persona = " ".join(persona.split())
            if not persona:
                continue
            clave = (normalizar(persona), clave_mesa(mesa), normalizar(jefe), rol)
            if clave not in vistos:
                vistos.add(clave)
                resultado.append({"nombre": persona, "mesa": mesa, "jefe": jefe, "rol": rol})
    for fila in filas[inicio+1:]:
        if not any(v is not None and str(v).strip() for v in fila):
            continue
        mesa = valor(fila, "mesa")
        jefe = valor(fila, "jefe")
        if normalizar(mesa) in alias["mesa"] and normalizar(jefe) in alias["jefe"]:
            continue  # encabezado repetido dentro de una tabla larga
        if mesa:
            if clave_mesa(mesa) != clave_mesa(mesa_actual):
                jefe_actual = ""  # nunca heredar el jefe de otra mesa
            mesa_actual = mesa
        if jefe:
            jefe_actual = jefe
        mesa, jefe = mesa_actual, jefe_actual
        if not mesa:
            sin_mesa += 1
            continue
        if jefe:
            agregar(jefe, "JEFE DE CRÉDITO", mesa, jefe)
        if indices["rol"] is not None and indices["nombre"] is not None:
            rol = valor(fila, "rol")
            nr = normalizar(rol)
            rol = ("ASESOR" if "asesor" in nr else "PROMOTOR" if "promotor" in nr
                   else "JEFE DE CRÉDITO" if "jefe" in nr else rol.upper() or "SIN ROL")
            agregar(valor(fila, "nombre"), rol, mesa, jefe)
        else:
            agregar(valor(fila, "asesor"), "ASESOR", mesa, jefe)
            agregar(valor(fila, "promotor"), "PROMOTOR", mesa, jefe)
        if not jefe:
            sin_jefe += 1
    if sin_mesa:
        avisos.append(f"Se omitieron {sin_mesa} filas sin mesa identificable.")
    if sin_jefe:
        avisos.append(f"Hay {sin_jefe} filas de estructura sin jefe asignado; revisa ESTRUCTURA DE COLABORADORES.")
    return resultado, avisos


def escoger_hoja(titulos, preferida):
    exactas = [t for t in titulos if normalizar(t) == normalizar(preferida)]
    if len(exactas) == 1:
        return exactas[0]
    alternativas = [t for t in titulos if normalizar(t) in {"colaboradores", "estructura de colaboradores"}]
    if len(alternativas) == 1:
        return alternativas[0]
    raise ValueError("ACCESOS debe contener la pestaña ESTRUCTURA DE COLABORADORES. Si tiene otro nombre, configúralo como hoja_colaboradores en [app] de Secrets.")


def descargar_excel(drive, archivo):
    limite = 20 * 1024 * 1024
    if int(archivo.get("size", 0)) > limite:
        raise ValueError("ACCESOS supera 20 MB. Usa una copia de menor tamaño o una hoja de Google.")
    buffer = io.BytesIO()
    peticion = drive.files().get_media(fileId=archivo["id"], supportsAllDrives=True)
    lector = MediaIoBaseDownload(buffer, peticion, chunksize=1024 * 1024)
    completo = False
    while not completo:
        _, completo = lector.next_chunk(num_retries=2)
        if buffer.tell() > limite:
            raise ValueError("ACCESOS supera 20 MB.")
    return buffer.getvalue()


def filas_excel(contenido, hoja):
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True, keep_links=False)
    try:
        titulo = escoger_hoja(libro.sheetnames, hoja)
        filas = [["" if v is None else v for v in fila] for fila in libro[titulo].iter_rows(values_only=True)]
        return filas
    finally:
        libro.close()


def localizar_accesos(drive, carpeta_id, archivo_id=""):
    if archivo_id:
        return drive.files().get(fileId=archivo_id, fields="id,name,mimeType,size",
                                 supportsAllDrives=True).execute(num_retries=2)
    carpeta_segura = carpeta_id.replace("\\", "\\\\").replace("'", "\\'")
    candidatos, token = [], None
    while True:
        respuesta = drive.files().list(
            q=f"'{carpeta_segura}' in parents and trashed = false",
            fields="nextPageToken,files(id,name,mimeType,size)", pageSize=1000,
            pageToken=token, spaces="drive", supportsAllDrives=True, includeItemsFromAllDrives=True,
        ).execute(num_retries=2)
        for f in respuesta.get("files", []):
            nombre = Path(f["name"]).stem if f["name"].lower().endswith((".xlsx", ".xls")) else f["name"]
            if normalizar(nombre) == "accesos" and (f["mimeType"] in {SHEET_MIME, XLSX_MIME}
                    or f["name"].lower().endswith((".xlsx", ".xls"))):
                candidatos.append(f)
        token = respuesta.get("nextPageToken")
        if not token:
            break
    if not candidatos:
        raise ValueError("No se encontró ACCESOS dentro de ESTRUCTURA Y ACCESOS. Comparte la carpeta como Lector con lector-colaboradores y comprueba el nombre del archivo.")
    if len(candidatos) > 1:
        raise ValueError("Hay varios archivos llamados ACCESOS. Configura accesos_id en [app] de Secrets con el ID del archivo que debe usarse.")
    return candidatos[0]


@st.cache_data(ttl=60, max_entries=8, show_spinner=False)
def leer_estructura(carpeta_id, archivo_id, hoja, credenciales_json):
    drive = servicio(credenciales_json)
    archivo = localizar_accesos(drive, carpeta_id, archivo_id)
    if archivo["mimeType"] == SHEET_MIME:
        sheets = servicio(credenciales_json, "sheets", "v4")
        metadata = sheets.spreadsheets().get(spreadsheetId=archivo["id"],
                    fields="sheets(properties(title))").execute(num_retries=2)
        titulo = escoger_hoja([s["properties"]["title"] for s in metadata.get("sheets", [])], hoja)
        filas = sheets.spreadsheets().values().get(spreadsheetId=archivo["id"],
            range="'"+titulo.replace("'", "''")+"'", valueRenderOption="FORMATTED_VALUE"
        ).execute(num_retries=2).get("values", [])
    elif archivo["mimeType"] == XLSX_MIME or archivo["name"].lower().endswith(".xlsx"):
        filas = filas_excel(descargar_excel(drive, archivo), hoja)
    else:
        raise ValueError("Guarda ACCESOS como .xlsx o conviértelo a una hoja de Google. El formato .xls no está admitido.")
    equipo, avisos = interpretar_estructura(filas)
    return equipo, avisos


def error_estructura(error):
    if isinstance(error, HttpError) and error.resp.status in (401, 403, 404):
        return "Comparte ESTRUCTURA Y ACCESOS y su archivo ACCESOS como Lector con la cuenta lector-colaboradores. Comprueba también Google Drive API y Google Sheets API."
    if isinstance(error, ValueError):
        return str(error)
    return "No se pudo leer ESTRUCTURA DE COLABORADORES. Revisa el archivo, su formato y los permisos de lectura."


def construir_equipo(fotos, estructura):
    """Conserva personas sin foto y fotos nuevas aún no asignadas a una mesa."""
    equipo, usadas = [], set()
    nombres_estructura = {normalizar(c["nombre"]) for c in estructura}
    for colaborador in estructura:
        nombre = normalizar(colaborador["nombre"])
        candidatas = [f for f in fotos if normalizar(f["colaborador"]) == nombre]
        if not candidatas:
            aproximadas = [f for f in fotos if coincide_abreviado(nombre, normalizar(f["colaborador"]))]
            if len(aproximadas) == 1:
                foto_nombre = normalizar(aproximadas[0]["colaborador"])
                if len({n for n in nombres_estructura if coincide_abreviado(n, foto_nombre)}) == 1:
                    candidatas = aproximadas
        foto = candidatas[0] if len(candidatas) == 1 else None
        if foto:
            usadas.add(foto["id"])
        c = dict(colaborador, foto=foto)
        clave = [normalizar(c["nombre"]), clave_mesa(c["mesa"]), normalizar(c["jefe"]), c["rol"]]
        c["id"] = "equipo_"+hashlib.sha256(json.dumps(clave).encode()).hexdigest()[:20]
        equipo.append(c)
    for foto in fotos:
        if foto["id"] not in usadas:
            equipo.append({"nombre": foto["colaborador"], "mesa": "", "jefe": "", "rol": "SIN ESTRUCTURA",
                           "foto": foto, "id": "foto_"+foto["id"]})
    return sorted(equipo, key=lambda c: (normalizar(c["mesa"]), normalizar(c["jefe"]), normalizar(c["nombre"]), c["rol"]))


def filtrar_equipo(equipo, mesa="Todas", jefe="Todos", rol="Todos", busqueda=""):
    return [c for c in equipo if (mesa == "Todas" or (mesa == "Sin mesa" and not c["mesa"])
                    or clave_mesa(c["mesa"]) == clave_mesa(mesa))
            and (jefe == "Todos" or (jefe == "Sin jefe asignado" and not c["jefe"])
                 or normalizar(c["jefe"]) == normalizar(jefe))
            and (rol == "Todos" or c["rol"] == rol)
            and normalizar(busqueda) in normalizar(c["nombre"])]


def agrupar_equipos(personas):
    """Una mesa es un equipo. El jefe es su responsable y no un subordinado."""
    grupos = {}
    for persona in personas:
        clave = (clave_mesa(persona["mesa"]), normalizar(persona["jefe"]))
        grupo = grupos.setdefault(clave, {"mesa": persona["mesa"], "jefe": persona["jefe"],
                                         "responsable": None, "integrantes": []})
        if persona["rol"] == "JEFE DE CRÉDITO":
            grupo["responsable"] = persona
        else:
            grupo["integrantes"].append(persona)
    return list(grupos.values())


def tabla_desempeno(integrantes, registros, nombres_equipo):
    """Una fila por integrante/asignación, sin sumar ni mezclar registros."""
    tabla = []
    for persona in integrantes:
        encontrados, tipo = buscar_desempeno(persona["nombre"], registros, nombres_equipo, persona["mesa"])
        registro = encontrados[0] if len(encontrados) == 1 else None
        estado = ("Disponible" if registro else "Varios registros: consultar tarjeta" if encontrados
                  else "Nombre ambiguo" if tipo == "ambigua" else "Sin datos")
        fila = {"Mesa / equipo": persona["mesa"] or "Sin mesa",
                "Jefe de crédito": persona["jefe"] or "Sin jefe asignado",
                "Colaborador": persona["nombre"], "Rol": persona["rol"], "Desempeño": estado}
        indicadores = registro["indicadores"] if registro else {}
        for etiqueta, aliases in INDICADORES_PRINCIPALES:
            fila[etiqueta] = next((v for k, v in indicadores.items() if normalizar(k) in aliases), "—")
        tabla.append(fila)
    return tabla


def nombre_colaborador(archivo):
    # Juan_Perez.jpg → Juan Perez. Conserva tildes y guiones de nombres compuestos.
    return " ".join(Path(archivo).stem.replace("_", " ").split())


def normalizar(texto):
    texto = str(texto).replace("_", " ").replace("-", " ")
    return " ".join("".join(c for c in unicodedata.normalize("NFD", texto.casefold())
                            if unicodedata.category(c) != "Mn").split())


LOGIN_ALIAS = {
    "usuario": {"usuario", "user", "username", "nombre de usuario", "login", "usuario de acceso", "correo", "email", "e mail"},
    "password": {"password", "contrasena", "clave", "pass", "contrasena de acceso", "password de acceso"},
    "activo": {"activo", "activa", "estado", "estatus", "status", "active", "habilitado", "habilitada", "estado de usuario", "acceso activo"},
    "nombre": {"personal", "nombre", "nombre completo", "nombre del usuario", "nombre y apellido", "colaborador", "persona"},
    "info_equipo": {"informacion del equipo", "informacion equipo", "info del equipo", "informacion de equipo", "acceso a informacion del equipo", "ver informacion del equipo"},
    "editar": {"editar", "permiso editar", "puede editar", "edicion"},
    "bonificacion": {"bonificacion", "bonificaciones", "permiso bonificacion", "ver bonificacion", "ver bonificaciones"},
    "rol": {"rol", "perfil", "tipo de usuario", "tipo usuario", "cargo", "nivel de acceso", "tipo de acceso", "rol de usuario"},
    "jefe": {"jefe de credito", "jefe de creditos", "jefe credito", "jefe", "nombre del jefe", "nombre jefe de credito", "jefe asignado"},
    "mesa": {"mesa", "mesas", "mesa asignada", "equipo", "equipo asignado"},
}

ESTADOS_ACTIVOS = {"activo", "activa", "si", "s", "1", "true", "yes", "habilitado", "habilitada", "vigente", "enabled"}
ROLES_ADMIN = {"admin", "administrador", "administradora", "superadmin", "administrador del sistema", "administradora del sistema"}
ROLES_JEFE = {"jefe", "jefe de credito", "jefe de creditos", "jefatura", "supervisor", "supervisora", "coordinador", "coordinadora", "manager", "gerente"}
ROLES_INDIVIDUALES = {"asesor", "asesora", "promotor", "promotora", "colaborador", "colaboradora", "personal"}


def estado_activo(valor):
    """Solo concede acceso cuando la celda contiene un estado afirmativo conocido."""
    if isinstance(valor, bool):
        return valor
    return normalizar(valor) in ESTADOS_ACTIVOS


def interpretar_login(filas):
    """Lee LOGIN sin depender del orden de columnas; por seguridad exige usuario, clave y estado."""
    indices, inicio = {}, None
    for i, fila in enumerate(filas[:10]):
        titulos = [normalizar(v) for v in fila]
        hallados = {campo: next((j for j, titulo in enumerate(titulos) if titulo in alias), None)
                    for campo, alias in LOGIN_ALIAS.items()}
        if all(hallados[campo] is not None for campo in ("usuario", "password", "activo")):
            indices, inicio = hallados, i
            break
    if inicio is None:
        raise ValueError("La pestaña LOGIN debe tener encabezados para USUARIO, PASSWORD/CONTRASEÑA y ACTIVO/ESTADO en las primeras 10 filas.")

    def valor(fila, campo, conservar_espacios=False):
        j = indices.get(campo)
        if j is None or j >= len(fila) or fila[j] is None:
            return ""
        resultado = str(fila[j])
        return resultado if conservar_espacios else resultado.strip()

    cuentas = []
    for fila in filas[inicio + 1:]:
        usuario = valor(fila, "usuario")
        if not usuario:
            continue
        cuentas.append({
            "usuario": usuario,
            "password": valor(fila, "password", conservar_espacios=True),
            "activo": estado_activo(valor(fila, "activo")),
            "nombre": valor(fila, "nombre"),
            "informacion_equipo": valor(fila, "info_equipo").strip().upper(),
            "editar": valor(fila, "editar").strip().upper(),
            "bonificacion": valor(fila, "bonificacion").strip().upper(),
            "rol": valor(fila, "rol"),
            "jefe": valor(fila, "jefe"),
            "mesa": valor(fila, "mesa"),
        })
    if not cuentas:
        raise ValueError("La pestaña LOGIN no contiene cuentas de usuario.")
    return cuentas


@st.cache_data(ttl=10, max_entries=8, show_spinner=False)
def leer_login(archivo_id, hoja, credenciales_json):
    """Lee LOGIN desde el mismo archivo ACCESOS, sea Google Sheets o .xlsx."""
    drive = servicio(credenciales_json)
    archivo = drive.files().get(
        fileId=archivo_id, fields="id,name,mimeType,size", supportsAllDrives=True
    ).execute(num_retries=2)
    if archivo["mimeType"] == SHEET_MIME:
        sheets = servicio(credenciales_json, "sheets", "v4")
        metadata = sheets.spreadsheets().get(
            spreadsheetId=archivo["id"], fields="sheets(properties(title))"
        ).execute(num_retries=2)
        titulos = [s["properties"]["title"] for s in metadata.get("sheets", [])]
        candidatas = [t for t in titulos if normalizar(t) == normalizar(hoja)]
        if len(candidatas) != 1:
            raise ValueError(f"No se encontró una pestaña única llamada {hoja} en ACCESOS.")
        rango = "'" + candidatas[0].replace("'", "''") + "'"
        filas = sheets.spreadsheets().values().get(
            spreadsheetId=archivo["id"], range=rango, majorDimension="ROWS",
            valueRenderOption="FORMATTED_VALUE",
        ).execute(num_retries=2).get("values", [])
    elif archivo["mimeType"] == XLSX_MIME or archivo["name"].lower().endswith(".xlsx"):
        filas = filas_excel(descargar_excel(drive, archivo), hoja)
    else:
        raise ValueError("ACCESOS debe ser una hoja de Google o un archivo .xlsx para poder leer LOGIN.")
    return interpretar_login(filas)


def buscar_cuenta_login(cuentas, usuario):
    """Devuelve una cuenta solo si el usuario coincide con una única fila."""
    buscado = normalizar(usuario)
    coincidencias = [c for c in cuentas if normalizar(c.get("usuario", "")) == buscado]
    return coincidencias[0] if buscado and len(coincidencias) == 1 else None


def password_correcta(ingresada, almacenada):
    """Admite clave actual en texto y SHA-256 explícito como sha256:<hex>."""
    ingresada, almacenada = str(ingresada), str(almacenada)
    if not almacenada:
        return False
    if almacenada.casefold().startswith("sha256:"):
        esperado = almacenada.split(":", 1)[1].strip().casefold()
        calculado = hashlib.sha256(ingresada.encode("utf-8")).hexdigest()
        return hmac.compare_digest(calculado, esperado)
    return hmac.compare_digest(ingresada, almacenada)


def cuentas_demo():
    return [{"usuario": "admin", "password": "demo123", "activo": True,
             "nombre": "Administrador de demostración", "informacion_equipo": "TOTAL",
             "editar": "SI", "bonificacion": "SI", "rol": "ADMINISTRADOR",
             "jefe": "", "mesa": ""}]


def limpiar_filtros_sesion():
    for clave in ("filtro_mesa", "filtro_jefe", "filtro_rol", "pagina"):
        st.session_state.pop(clave, None)


def cerrar_sesion():
    st.session_state.pop("login_usuario", None)
    limpiar_filtros_sesion()


def mostrar_login(cuentas):
    """Muestra el formulario y devuelve la cuenta vigente; nunca guarda la contraseña en sesión."""
    usuario_sesion = st.session_state.get("login_usuario", "")
    if usuario_sesion:
        cuenta = buscar_cuenta_login(cuentas, usuario_sesion)
        if cuenta and cuenta.get("activo"):
            return cuenta
        cerrar_sesion()
        st.warning("Tu cuenta está inactiva, ya no existe o está duplicada en LOGIN. Ingresa con una cuenta autorizada.")

    st.title("🔐 Acceso a desempeño")
    st.caption("Ingresa con las credenciales registradas en la pestaña LOGIN de ACCESOS.")
    _, centro, _ = st.columns([1, 1.2, 1])
    with centro:
        with st.form("formulario_login", clear_on_submit=False):
            usuario = st.text_input("Usuario", placeholder="Escribe tu usuario", autocomplete="username")
            password = st.text_input("Contraseña", type="password", placeholder="Escribe tu contraseña", autocomplete="current-password")
            ingresar = st.form_submit_button("Ingresar", type="primary", use_container_width=True)
        if ingresar:
            cuenta = buscar_cuenta_login(cuentas, usuario)
            if cuenta and cuenta.get("activo") and password_correcta(password, cuenta.get("password", "")):
                limpiar_filtros_sesion()
                st.session_state["login_usuario"] = cuenta["usuario"]
                st.rerun()
            st.error("No se pudo ingresar. Verifica tus credenciales y consulta si tu cuenta está activa.")
    return None


def aplicar_permisos_equipo(equipo, cuenta):
    """
    El alcance se determina combinando LOGIN con ESTRUCTURA DE COLABORADORES:
    TOTAL = todos los equipos; SI = equipo completo solo si PERSONAL figura como
    JEFE DE CRÉDITO; SI para asesores/promotores = solo su propia ficha;
    NO = solo su propia ficha, independientemente del cargo.
    """
    permiso = normalizar(cuenta.get("informacion_equipo", ""))
    nombre = str(cuenta.get("nombre", "") or "").strip()
    usuario = str(cuenta.get("usuario", "") or "").strip()
    identidad = normalizar(nombre or usuario)

    if permiso == "total":
        return equipo, "Acceso total a todos los equipos", "total", ""

    if permiso not in {"si", "no"}:
        return [], "", "", (
            "La cuenta está activa, pero INFORMACIÓN DEL EQUIPO debe indicar SI, NO o TOTAL. "
            "Revisa esa columna en la pestaña LOGIN."
        )

    coincidencias = [
        c for c in equipo
        if identidad and normalizar(c.get("nombre", "")) == identidad
        and c.get("rol") != "SIN ESTRUCTURA"
    ]
    if not coincidencias:
        return [], "", "", (
            "El valor de PERSONAL en LOGIN no coincide con un nombre de "
            "ESTRUCTURA DE COLABORADORES. Revisa que esté escrito igual."
        )

    # NO siempre limita al registro propio, incluso si la persona es jefe.
    if permiso == "no":
        return coincidencias, f"Acceso individual: {nombre or usuario}", "individual", ""

    # SI no convierte a un asesor/promotor en jefe. La estructura decide el rol.
    jefaturas = [
        c for c in coincidencias
        if normalizar(c.get("rol", "")) == normalizar("JEFE DE CRÉDITO")
        and c.get("mesa") and c.get("jefe")
    ]
    if not jefaturas:
        propios = [
            c for c in coincidencias
            if normalizar(c.get("rol", "")) in {
                normalizar("ASESOR"), normalizar("PROMOTOR"),
                normalizar("COLABORADOR"), normalizar("PERSONAL")
            }
        ]
        if not propios:
            # Para cualquier rol que no sea jefe, la opción segura es no ampliar
            # el acceso al equipo por el mero hecho de que LOGIN diga SI.
            propios = coincidencias
        return propios, f"Acceso individual: {nombre or usuario}", "individual", ""

    # Solo las mesas en las que la estructura registra a esta persona como jefe.
    asignaciones_jefe = {
        (clave_mesa(c.get("mesa", "")), normalizar(c.get("jefe", "")))
        for c in jefaturas
    }
    autorizados = [
        c for c in equipo
        if (clave_mesa(c.get("mesa", "")), normalizar(c.get("jefe", ""))) in asignaciones_jefe
    ]
    if not autorizados:
        return [], "", "", "No se encontraron colaboradores para las mesas asignadas a este jefe."

    mesas = sorted({c.get("mesa", "") for c in jefaturas if c.get("mesa")})
    descripcion = f"Equipo a cargo de {nombre or usuario}: " + ", ".join(mesas)
    return autorizados, descripcion, "equipo", ""


def error_login(error):
    if isinstance(error, HttpError) and error.resp.status in (401, 403, 404):
        return "No se puede leer ACCESOS. Comparte el archivo con la cuenta de servicio como Lector y revisa el ID."
    if isinstance(error, ValueError):
        return str(error)
    return "No se pudo verificar la pestaña LOGIN. Revisa la conexión, los permisos y vuelve a intentarlo."


def servicio(credenciales_json, api="drive", version="v3"):
    credenciales = Credentials.from_service_account_info(
        json.loads(credenciales_json), scopes=SCOPE
    )
    return build(api, version, credentials=credenciales, cache_discovery=False)


def interpretar_desempeno(filas):
    """Detecta el encabezado y conserva los valores tal como se ven en Sheets."""
    if not filas:
        return [], []
    inicio = None
    for numero, fila in enumerate(filas[:10]):
        if any(normalizar(v) in NOMBRES_ASESOR for v in fila):
            inicio = numero
            break
    if inicio is None:
        raise ValueError("No se encontró una columna ASESOR, COLABORADOR, PERSONAL o NOMBRE en las primeras 10 filas de DESEMPEÑO.")
    cabecera = [str(v).strip() for v in filas[inicio]]
    avisos = []
    # El script original explica que A está mal rotulada CLIENTES NUEVO,
    # mientras que B es ASESOR y los indicadores ocupan C:M.
    if (len(cabecera) >= 13 and normalizar(cabecera[0]) == "clientes nuevo"
            and normalizar(cabecera[1]) in NOMBRES_ASESOR):
        cabecera[:13] = ORDEN_ORIGINAL
        avisos.append("Se interpretó A:M con el orden de tu script original porque la primera columna está rotulada CLIENTES NUEVO y contiene la mesa.")
    columna_nombre = next(i for i, v in enumerate(cabecera) if normalizar(v) in NOMBRES_ASESOR)
    columna_mesa = next((i for i, v in enumerate(cabecera) if normalizar(v) in {"mesa", "mesas"}), None)
    resultado = []
    for indice, fila in enumerate(filas[inicio+1:], start=inicio+2):
        nombre = str(fila[columna_nombre]).strip() if columna_nombre < len(fila) else ""
        if not nombre:
            continue
        mesa = str(fila[columna_mesa]).strip() if columna_mesa is not None and columna_mesa < len(fila) else ""
        indicadores, repetidos = {}, {}
        for i, titulo in enumerate(cabecera):
            if i in (columna_nombre, columna_mesa) or not titulo:
                continue
            clave = normalizar(titulo)
            repetidos[clave] = repetidos.get(clave, 0)+1
            etiqueta = titulo if repetidos[clave] == 1 else f"{titulo} ({repetidos[clave]})"
            valor = fila[i] if i < len(fila) else ""
            indicadores[etiqueta] = "—" if valor is None or str(valor).strip() == "" else str(valor)
        resultado.append({"nombre": nombre, "mesa": mesa, "fila": indice, "indicadores": indicadores})
    return resultado, avisos


def coincide_abreviado(a, b):
    return a == b or a.startswith(b + " ") or b.startswith(a + " ")


def buscar_desempeno(nombre, registros, nombres_fotos, mesa=""):
    # La misma persona puede aparecer en varias mesas. Nunca tomar el desempeño
    # de otra mesa; si la hoja no trae mesa, se permite la coincidencia por nombre.
    if mesa:
        misma_mesa = [r for r in registros if clave_mesa(r["mesa"]) == clave_mesa(mesa)]
        registros = misma_mesa + [r for r in registros if not r["mesa"]]
    buscado = normalizar(nombre)
    exactos = [r for r in registros if normalizar(r["nombre"]) == buscado]
    if exactos:
        return exactos, "exacta"
    candidatos = [r for r in registros if coincide_abreviado(buscado, normalizar(r["nombre"]))]
    nombres = {normalizar(r["nombre"]) for r in candidatos}
    if len(nombres) != 1:
        return [], "ambigua" if nombres else "sin datos"
    nombre_origen = next(iter(nombres))
    posibles_fotos = {normalizar(n) for n in nombres_fotos if coincide_abreviado(normalizar(n), nombre_origen)}
    if len(posibles_fotos) != 1:
        return [], "ambigua"
    return candidatos, "abreviada"


@st.cache_data(ttl=60, max_entries=8, show_spinner=False)
def leer_desempeno(archivo_id, hoja, credenciales_json):
    sheets = servicio(credenciales_json, "sheets", "v4")
    metadata = sheets.spreadsheets().get(
        spreadsheetId=archivo_id, fields="sheets(properties(title))"
    ).execute(num_retries=2)
    titulos = [s["properties"]["title"] for s in metadata.get("sheets", [])]
    candidatas = [t for t in titulos if normalizar(t) == normalizar(hoja)]
    if len(candidatas) != 1:
        raise ValueError("No se encontró una pestaña única llamada " + hoja +
                         ". Revisa su nombre; puede configurarse como hoja_desempeno en la sección [app] de Secrets.")
    rango = "'" + candidatas[0].replace("'", "''") + "'"
    respuesta = sheets.spreadsheets().values().get(
        spreadsheetId=archivo_id, range=rango, majorDimension="ROWS",
        valueRenderOption="FORMATTED_VALUE",
    ).execute(num_retries=2)
    return interpretar_desempeno(respuesta.get("values", []))


def error_productividad(error):
    if isinstance(error, HttpError):
        if error.resp.status == 404:
            return "Comparte el archivo PRODUCTIVIDAD como Lector con la cuenta lector-colaboradores y revisa su ID."
        if error.resp.status in (401, 403):
            return "Habilita Google Sheets API en el mismo proyecto y comparte PRODUCTIVIDAD como Lector con la cuenta lector-colaboradores."
        if error.resp.status == 429:
            return "Se alcanzó la cuota de Sheets. Espera un minuto y pulsa Actualizar ahora."
        if error.resp.status == 400:
            return "El archivo debe ser una Hoja de cálculo de Google y tener la pestaña DESEMPEÑO. Si es un Excel, conviértelo a Google Sheets y configura el ID de la copia en productividad_id."
    return str(error) if isinstance(error, ValueError) else "No se pudo consultar PRODUCTIVIDAD. Revisa la conexión y vuelve a actualizar."


def numero_indicador(valor):
    """Convierte porcentajes y cantidades visibles a número para aplicar semáforos."""
    texto = str(valor or "").strip().replace("S/", "").replace("%", "").replace(" ", "")
    if not texto or texto in {"—", "-", "N/A", "NA"}:
        return None
    # Conserva signo y decimal. En formatos 1,234.50 elimina separadores de miles.
    if "," in texto and "." in texto:
        texto = texto.replace(",", "")
    elif "," in texto:
        trozos = texto.split(",")
        if len(trozos) == 2 and len(trozos[1]) in (1, 2):
            texto = ".".join(trozos)
        else:
            texto = "".join(trozos)
    coincidencia = re.search(r"[-+]?\d+(?:\.\d+)?", texto)
    if not coincidencia:
        return None
    try:
        return float(coincidencia.group(0))
    except ValueError:
        return None


def estilo_indicador(nombre, valor):
    """Devuelve fondo, borde, texto y estado según las metas acordadas."""
    clave = normalizar(nombre)
    numero = numero_indicador(valor)
    verde = ("#DCFCE7", "#22C55E", "#166534")
    naranja = ("#FFEDD5", "#F97316", "#9A3412")
    rojo = ("#FEE2E2", "#EF4444", "#991B1B")
    rojo_fuerte = ("#FECACA", "#DC2626", "#7F1D1D")
    rojo_oscuro = ("#FCA5A5", "#B91C1C", "#7F1D1D")
    amarillo = ("#FEF3C7", "#F59E0B", "#92400E")
    neutro = ("#F8FAFC", "#CBD5E1", "#334155")

    if numero is None:
        return (*neutro, "Sin dato")
    if "variacion" in clave and "cliente" in clave:
        return (*verde, "Variación positiva") if numero > 0 else (*rojo, "Variación no positiva")
    if "bolsas medio dia" in clave or "bolsas mediodia" in clave:
        if numero >= 69:
            return (*verde, "Meta alcanzada · ≥ 69%")
        if numero > 64:
            return (*naranja, "Atención · 65–68.99%")
        return (*rojo, "Bajo mínimo · ≤ 64%")
    es_cierre = ("cierre bolsa" in clave or "bolsas cierre" in clave or
                 ("bolsas" in clave and "cierre" in clave))
    if es_cierre:
        if numero >= 95:
            return (*verde, "Meta alcanzada · ≥ 95%")
        if numero >= 90:
            return (*amarillo, "Cerca de la meta · 90–94.99%")
        if numero >= 80:
            return (*naranja, "Requiere mejora · 80–89.99%")
        if numero >= 70:
            return (*rojo, "Bajo · 70–79.99%")
        return (*rojo_oscuro, "Crítico · < 70%")
    return (*neutro, "Indicador")


def prioridad_indicador(nombre):
    clave = normalizar(nombre)
    if "bolsas medio dia" in clave or "bolsas mediodia" in clave:
        return 0
    if "cierre bolsa" in clave or "bolsas cierre" in clave or ("bolsas" in clave and "cierre" in clave):
        return 1
    if "variacion" in clave and "cliente" in clave:
        return 2
    if "clientes nuevo" in clave:
        return 3
    if "clientes recuperado" in clave:
        return 4
    return 10


def inyectar_estilos_dashboard():
    st.markdown("""
    <style>
      .kpi-grid { display:grid; grid-template-columns:repeat(auto-fit,minmax(185px,1fr)); gap:12px; width:100%; margin:10px 0 20px; }
      .kpi-card { border:1px solid; border-radius:16px; padding:16px 17px 14px; min-height:112px; box-sizing:border-box; box-shadow:0 2px 8px rgba(15,23,42,.045); }
      .kpi-label { font-size:.79rem; font-weight:650; letter-spacing:.025em; text-transform:uppercase; line-height:1.35; opacity:.9; overflow-wrap:anywhere; }
      .kpi-value { font-size:1.8rem; font-weight:750; line-height:1.25; margin-top:9px; overflow-wrap:anywhere; }
      .kpi-status { font-size:.72rem; margin-top:8px; line-height:1.3; opacity:.9; }
      .person-card { border:1px solid #E2E8F0; border-radius:18px; background:#FFFFFF; padding:14px; text-align:center; margin-bottom:12px; box-shadow:0 3px 12px rgba(15,23,42,.06); min-height:260px; }
      .person-photo-link { display:block; width:100%; max-width:205px; margin:0 auto 12px; border-radius:14px; overflow:hidden; border:2px solid transparent; }
      .person-photo-link:hover { border-color:#14B8A6; }
      .person-photo-link img { display:block; width:100%; aspect-ratio:1/1; object-fit:cover; transition:transform .18s ease; }
      .person-photo-link:hover img { transform:scale(1.035); }
      .person-name { color:#0F172A; font-size:1rem; font-weight:750; line-height:1.3; margin:7px 0 5px; overflow-wrap:anywhere; }
      .person-meta { color:#64748B; font-size:.82rem; margin-bottom:8px; }
      .person-open { color:#0F766E; font-size:.8rem; font-weight:700; text-decoration:none; }
      .team-banner { border-radius:14px; padding:12px 16px; margin:16px 0 10px; background:linear-gradient(110deg,#F0FDFA,#F8FAFC); border:1px solid #CCFBF1; color:#134E4A; }
      .profile-header { border:1px solid #E2E8F0; border-radius:20px; padding:22px; text-align:center; background:linear-gradient(160deg,#FFFFFF,#F8FAFC); box-shadow:0 3px 14px rgba(15,23,42,.05); margin:8px 0 20px; }
      .profile-header img { width:128px; height:128px; object-fit:cover; border-radius:22px; border:3px solid #CCFBF1; }
      .profile-title { color:#0F172A; font-size:1.65rem; font-weight:800; margin:12px 0 4px; line-height:1.2; }
      .profile-subtitle { color:#64748B; margin:0 0 8px; font-size:.95rem; }
      .profile-chip { display:inline-block; background:#F1F5F9; color:#334155; border:1px solid #E2E8F0; border-radius:999px; padding:5px 10px; font-size:.8rem; margin:3px; }
      @media (max-width:600px) { .kpi-grid { grid-template-columns:repeat(2,minmax(0,1fr)); gap:8px; } .kpi-card { padding:12px; min-height:100px; } .kpi-value { font-size:1.35rem; } }
    </style>
    """, unsafe_allow_html=True)


def mostrar_indicadores(registros, tipo, archivo_id):
    """Muestra los indicadores como tarjetas de dashboard, ordenadas y con semáforos."""
    if not registros:
        st.caption("Nombre ambiguo: haz coincidir el nombre completo de la foto y el de la hoja." if tipo == "ambigua"
                   else "Sin desempeño asociado a esta persona y mesa. Revisa el nombre en la hoja.")
        return
    if len(registros) > 1:
        st.caption("Hay varios registros para este nombre; selecciona el que corresponde a la mesa.")
        indice = st.selectbox("Registro de desempeño", list(range(len(registros))), index=None,
                              placeholder="Selecciona una mesa o fila",
                              format_func=lambda i: f"{registros[i]['mesa'] or 'Sin mesa'} · fila {registros[i]['fila']}",
                              key="registro_" + archivo_id)
        if indice is None:
            return
        registro = registros[indice]
    else:
        registro = registros[0]
    if tipo == "abreviada":
        st.caption("Coincidencia por nombre abreviado: " + registro["nombre"])
    if registro.get("mesa"):
        st.caption("Mesa de desempeño: " + str(registro["mesa"]))
    indicadores = registro.get("indicadores", {})
    if not indicadores:
        st.info("La fila tiene nombre, pero no tiene columnas de indicadores.")
        return
    tarjetas = []
    for etiqueta, valor in sorted(indicadores.items(), key=lambda item: (prioridad_indicador(item[0]), normalizar(item[0]))):
        fondo, borde, texto, estado = estilo_indicador(etiqueta, valor)
        tarjetas.append(
            '<div class="kpi-card" style="background:' + fondo + ';border-color:' + borde + ';color:' + texto + ';">'
            '<div class="kpi-label">' + html.escape(str(etiqueta)) + '</div>'
            '<div class="kpi-value">' + html.escape(str(valor if str(valor).strip() else "—")) + '</div>'
            '<div class="kpi-status">' + html.escape(estado) + '</div></div>'
        )
    st.markdown('<div class="kpi-grid">' + ''.join(tarjetas) + '</div>', unsafe_allow_html=True)


@st.cache_data(ttl=60, max_entries=8, show_spinner=False)
def listar_fotos(carpeta_id, credenciales_json):
    drive = servicio(credenciales_json)
    carpeta = drive.files().get(
        fileId=carpeta_id, fields="name,mimeType", supportsAllDrives=True
    ).execute(num_retries=2)
    if carpeta.get("mimeType") != "application/vnd.google-apps.folder":
        raise ValueError("El ID configurado no corresponde a una carpeta de Google Drive.")
    fotos, omitidas, token = [], 0, None
    carpeta_segura = carpeta_id.replace("\\", "\\\\").replace("'", "\\'")
    while True:
        respuesta = drive.files().list(
            q=f"'{carpeta_segura}' in parents and trashed = false and mimeType contains 'image/'",
            fields="nextPageToken,files(id,name,mimeType,size,modifiedTime)",
            pageSize=1000, pageToken=token, spaces="drive",
            supportsAllDrives=True, includeItemsFromAllDrives=True,
        ).execute(num_retries=2)
        for archivo in respuesta.get("files", []):
            if archivo["mimeType"] not in FORMATOS:
                omitidas += 1
                continue
            archivo["colaborador"] = nombre_colaborador(archivo["name"])
            fotos.append(archivo)
        token = respuesta.get("nextPageToken")
        if not token:
            break
    return sorted(fotos, key=lambda f: (normalizar(f["colaborador"]), f["id"])), omitidas


@st.cache_data(ttl=3600, max_entries=300, show_spinner=False)
def descargar_foto(archivo_id, modificada, credenciales_json):
    # modificada forma parte de la clave de caché para detectar fotos reemplazadas.
    drive = servicio(credenciales_json)
    peticion = drive.files().get_media(fileId=archivo_id, supportsAllDrives=True)
    salida = io.BytesIO()
    descarga = MediaIoBaseDownload(salida, peticion, chunksize=1024 * 1024)
    terminado = False
    while not terminado:
        _, terminado = descarga.next_chunk(num_retries=2)
        if salida.tell() > MAX_BYTES:
            raise ValueError("La foto supera 10 MB. Sube una versión más pequeña.")
    salida.seek(0)
    with Image.open(salida) as original:
        imagen = ImageOps.exif_transpose(original).convert("RGB")
        imagen = ImageOps.fit(imagen, (450, 450), method=Image.Resampling.LANCZOS)
        miniatura = io.BytesIO()
        imagen.save(miniatura, format="JPEG", quality=85)
        return miniatura.getvalue()


def error_amigable(error):
    if isinstance(error, HttpError):
        if error.resp.status == 404:
            return "No se puede acceder a la carpeta o foto. Revisa su ID y compártela con el correo de la cuenta de servicio."
        if error.resp.status in (401, 403):
            return "Google denegó el acceso. Habilita Google Drive API y revisa los permisos de la cuenta de servicio."
        if error.resp.status == 429:
            return "Se alcanzó la cuota de Google. Espera un minuto y vuelve a intentar."
    if isinstance(error, ValueError):
        return str(error)
    return "No se pudo leer Google Drive. Revisa la conexión y la configuración."


def configuracion():
    try:
        valores = st.secrets.to_dict()
    except FileNotFoundError:
        valores = {}
    opciones = valores.get("app", {})
    return (str(opciones.get("folder_id", CARPETA_PREDETERMINADA)),
            valores.get("gcp_service_account"),
            bool(opciones.get("demo_mode", False)),
            str(opciones.get("productividad_id", PRODUCTIVIDAD_PREDETERMINADA)),
            str(opciones.get("hoja_desempeno", "DESEMPEÑO")),
            str(opciones.get("estructura_carpeta_id", ESTRUCTURA_CARPETA_PREDETERMINADA)),
            str(opciones.get("accesos_id", ACCESOS_PREDETERMINADO)),
            str(opciones.get("hoja_colaboradores", "ESTRUCTURA DE COLABORADORES")),
            str(opciones.get("hoja_login", LOGIN_HOJA_PREDETERMINADA)))


def foto_demo(indice):
    colores = ["#087E8B", "#59768B", "#AE647D"]
    img = Image.new("RGB", (450, 450), colores[indice % len(colores)])
    # Silueta sencilla generada por código: datos de prueba, sin fotos reales.
    from PIL import ImageDraw
    draw = ImageDraw.Draw(img)
    draw.ellipse((155, 70, 295, 210), fill="#EAF0F3")
    draw.rounded_rectangle((85, 235, 365, 485), radius=100, fill="#EAF0F3")
    buffer = io.BytesIO()
    img.save(buffer, "PNG")
    return buffer.getvalue()


def opciones_unicas(valores, clave=normalizar):
    opciones = {}
    for valor in valores:
        if valor:
            opciones.setdefault(clave(valor), valor)
    return sorted(opciones.values(), key=normalizar)


def filtro(etiqueta, opciones, key):
    # Si el usuario cambia de mesa o se actualiza Drive, la selección anterior
    # puede dejar de existir entre las opciones dependientes.
    if st.session_state.get(key) not in opciones:
        st.session_state[key] = opciones[0]
    return st.selectbox(etiqueta, opciones, key=key)


def obtener_foto_bytes(persona, demo, credenciales_json):
    """Devuelve bytes, MIME y si se trata de una foto pendiente."""
    archivo = persona.get("foto")
    if not archivo:
        return foto_demo(1), "image/png", True
    if demo:
        return foto_demo(int(archivo["id"])), "image/png", False
    if int(archivo.get("size", 0)) > MAX_BYTES:
        raise ValueError("La foto supera 10 MB.")
    return descargar_foto(archivo["id"], archivo.get("modifiedTime", ""), credenciales_json), "image/jpeg", False


def mostrar_foto(persona, demo, credenciales_json, ancho="stretch"):
    try:
        contenido, _, pendiente = obtener_foto_bytes(persona, demo, credenciales_json)
        st.image(contenido, width=ancho)
        if pendiente:
            st.caption("Foto pendiente")
    except Exception:
        st.warning("No se pudo cargar esta foto. Comprueba el formato, tamaño y permisos.")


def mostrar_tarjeta_colaborador(persona, demo, credenciales_json):
    """La foto es el enlace que abre la ficha individual; la tarjeta no revela KPI."""
    try:
        contenido, mime, pendiente = obtener_foto_bytes(persona, demo, credenciales_json)
    except Exception:
        contenido, mime, pendiente = foto_demo(1), "image/png", False
    imagen = base64.b64encode(contenido).decode("ascii")
    url = "?colaborador=" + quote(str(persona["id"]))
    nombre = html.escape(str(persona.get("nombre", "Sin nombre")))
    rol = html.escape(str(persona.get("rol", "" )).title())
    mesa = html.escape(str(persona.get("mesa", "Sin mesa") or "Sin mesa"))
    nota = '<div class="person-meta">Foto pendiente</div>' if pendiente else ""
    st.markdown(
        '<div class="person-card">'
        '<a class="person-photo-link" href="' + url + '" target="_self" title="Abrir desempeño de ' + nombre + '">'
        '<img src="data:' + mime + ';base64,' + imagen + '" alt="Foto de ' + nombre + '"></a>'
        '<div class="person-name">' + nombre + '</div>'
        '<div class="person-meta">' + rol + ' · ' + mesa + '</div>' + nota +
        '<a class="person-open" href="' + url + '" target="_self">Ver desempeño →</a>'
        '</div>', unsafe_allow_html=True
    )


def mostrar_perfil_dashboard(persona, demo, credenciales_json):
    """Cabecera vertical de perfil, seguida de las tarjetas de indicadores."""
    try:
        contenido, mime, pendiente = obtener_foto_bytes(persona, demo, credenciales_json)
    except Exception:
        contenido, mime, pendiente = foto_demo(1), "image/png", False
    imagen = base64.b64encode(contenido).decode("ascii")
    nombre = html.escape(str(persona.get("nombre", "Sin nombre")))
    rol = html.escape(str(persona.get("rol", "")).title())
    mesa = html.escape(str(persona.get("mesa", "Sin mesa") or "Sin mesa"))
    jefe = html.escape(str(persona.get("jefe", "") or ""))
    chip_jefe = '<span class="profile-chip">Jefe de crédito: ' + jefe + '</span>' if jefe and normalizar(rol) != normalizar("Jefe De Crédito") else ""
    chip_foto = '<span class="profile-chip">Foto pendiente</span>' if pendiente else ""
    st.markdown(
        '<div class="profile-header"><img src="data:' + mime + ';base64,' + imagen + '" alt="Foto de ' + nombre + '">'
        '<div class="profile-title">' + nombre + '</div><div class="profile-subtitle">' + rol + '</div>'
        '<span class="profile-chip">Mesa: ' + mesa + '</span>' + chip_jefe + chip_foto + '</div>',
        unsafe_allow_html=True
    )


def main():
    st.set_page_config(page_title="Desempeño · Colaboradores", page_icon="👥", layout="wide")
    try:
        (carpeta_id, credenciales, demo, productividad_id, hoja_desempeno,
         estructura_carpeta_id, accesos_id, hoja_colaboradores, hoja_login) = configuracion()
    except Exception:
        st.error("La configuración no es válida. Revisa Settings → Secrets.")
        return
    if not demo and not credenciales:
        st.title("🔐 Acceso a desempeño")
        st.info("Falta conectar Google Drive para validar los usuarios de la pestaña LOGIN.")
        st.write("**Para conectar la aplicación:**")
        st.write("1. Habilita Google Drive API y Google Sheets API en Google Cloud.")
        st.write("2. Comparte ACCESOS, PRODUCTIVIDAD y la carpeta de fotos con la cuenta de servicio, como Lector.")
        st.write("3. Coloca las credenciales en Settings → Secrets de Streamlit.")
        st.caption("No subas la clave privada a GitHub.")
        return
    cred_json = json.dumps(credenciales or {}, sort_keys=True)
    if demo:
        cuentas = cuentas_demo()
    else:
        try:
            cuentas = leer_login(accesos_id, hoja_login, cred_json)
        except Exception as error:
            st.title("🔐 Acceso a desempeño")
            st.error("No se puede mostrar el formulario porque no fue posible validar LOGIN.")
            st.caption(error_login(error))
            return
    cuenta = mostrar_login(cuentas)
    if not cuenta:
        return

    inyectar_estilos_dashboard()
    st.title("📊 Desempeño")
    st.caption("Panel de desempeño por equipos. Las fichas y los indicadores respetan los permisos de tu cuenta.")
    with st.sidebar:
        st.subheader("Sesión")
        st.caption(f"Usuario: {cuenta.get('usuario', '')}")
        st.caption(f"Información del equipo: {cuenta.get('informacion_equipo', 'NO') or 'NO'}")
        if st.button("Cerrar sesión", use_container_width=True):
            cerrar_sesion()
            st.query_params.clear()
            st.rerun()
        st.divider()
        st.subheader("Directorio")
        if st.button("Actualizar ahora"):
            listar_fotos.clear()
            leer_desempeno.clear()
            leer_estructura.clear()
            leer_login.clear()
        automatico = st.checkbox("Actualizar cada minuto", value=True)
        st.caption("La actualización funciona mientras la pestaña esté abierta y activa. El estado de LOGIN se valida periódicamente.")
        st.caption("Las fotos y la estructura se consultan desde Google Drive.")

    @st.fragment(run_every=15 if automatico else None)
    def galeria():
        # Revalida la cuenta activa antes de renderizar cualquier información.
        if demo:
            cuentas_actuales = cuentas_demo()
        else:
            try:
                cuentas_actuales = leer_login(accesos_id, hoja_login, cred_json)
            except Exception:
                cerrar_sesion()
                st.query_params.clear()
                st.error("No se pudo verificar tu acceso. Por seguridad, la sesión se cerró.")
                st.rerun()
                return
        cuenta_actual = buscar_cuenta_login(cuentas_actuales, st.session_state.get("login_usuario", ""))
        if not cuenta_actual or not cuenta_actual.get("activo"):
            cerrar_sesion()
            st.query_params.clear()
            st.warning("Tu cuenta fue desactivada o dejó de estar autorizada. Inicia sesión nuevamente.")
            st.rerun()
            return

        if demo:
            st.info("DEMOSTRACIÓN · Nombres e imágenes ficticios; no se está leyendo Drive.")
            fotos = [{"id": str(i), "colaborador": n} for i, n in enumerate(
                ["Ana Ejemplo", "Carlos Ejemplo", "María Ejemplo"])]
            omitidas = 0
            registros, avisos = interpretar_desempeno([
                ["MESA", "ASESOR", "BOLSAS MEDIO DIA", "BOLSAS CIERRE", "CLIENTES RECUPERADO", "CLIENTES NUEVO", "VARIACION CLIENTES", "NO RENOVADOS", "VARIACION DE CARTERA", "RECUPERACIÓN DE EFECTIVO", "RECUPERACIÓN DE MORA (30%)", "RECUPERACION MORA AGOSTO", "MORA SEPTIEMBRE"],
                ["MESA 1", "Ana Ejemplo", "69%", "96%", "3", "4", "2", "1", "S/ 1,000.00", "S/ 700.00", "30%", "S/ 100.00", "S/ 200.00"],
                ["MESA 1", "Carlos Ejemplo", "66%", "88%", "4", "2", "0", "2", "S/ 1,500.00", "S/ 900.00", "25%", "S/ 150.00", "S/ 250.00"],
                ["MESA 2", "María Ejemplo", "62%", "73%", "5", "3", "-2", "0", "S/ 2,000.00", "S/ 800.00", "40%", "S/ 120.00", "S/ 300.00"],
            ])
            estructura, avisos_estructura = interpretar_estructura([
                ["MESAS", "JEFE DE CREDITO", "ASESORES", "PROMOTORES"],
                ["MESA 1", "Lucía Ejemplo", "Ana Ejemplo", "Luis Ejemplo"],
                ["", "", "Carlos Ejemplo", ""],
                ["MESA 2", "Pedro Ejemplo", "María Ejemplo", ""],
            ])
            fallo_desempeno = False
        else:
            try:
                fotos, omitidas = listar_fotos(carpeta_id, cred_json)
            except Exception as error:
                fotos, omitidas = [], 0
                st.warning("No se pudieron consultar las fotos. " + error_amigable(error))
            try:
                registros, avisos = leer_desempeno(productividad_id, hoja_desempeno, cred_json)
                fallo_desempeno = False
            except Exception as error:
                registros, avisos = [], []
                fallo_desempeno = True
                st.warning("Desempeño no disponible. " + error_productividad(error))
            try:
                estructura, avisos_estructura = leer_estructura(
                    estructura_carpeta_id, accesos_id, hoja_colaboradores, cred_json)
            except Exception as error:
                estructura, avisos_estructura = [], []
                st.warning("Estructura no disponible. " + error_estructura(error))

        for aviso in avisos + avisos_estructura:
            st.info(aviso)
        if omitidas:
            st.warning(f"Hay {omitidas} imágenes con formato no compatible. Usa JPG, PNG, WEBP o GIF.")
        equipo_completo = construir_equipo(fotos, estructura)
        if not equipo_completo:
            st.info("Todavía no hay colaboradores para mostrar. Revisa los avisos, añade personas a ESTRUCTURA DE COLABORADORES o coloca fotos con sus nombres en la carpeta de fotos.")
            return
        nombres_equipo_completo = [c["nombre"] for c in equipo_completo]
        equipo, descripcion_acceso, tipo_acceso, error_acceso = aplicar_permisos_equipo(equipo_completo, cuenta_actual)
        if not equipo:
            st.error(error_acceso or "No hay información asignada a esta cuenta.")
            st.caption("Un administrador debe revisar el usuario en LOGIN y su relación con ESTRUCTURA DE COLABORADORES.")
            return
        st.caption("Acceso autorizado · " + descripcion_acceso)

        # Asesores, promotores y cuentas limitadas llegan directamente a su ficha propia.
        if tipo_acceso == "individual":
            st.subheader("Mi dashboard de desempeño")
            st.caption("Esta cuenta solo puede consultar su propia información.")
            for persona in equipo:
                mostrar_perfil_dashboard(persona, demo, cred_json)
                if fallo_desempeno:
                    st.warning("Tu desempeño no está disponible en este momento. Revisa el aviso de conexión.")
                else:
                    registros_persona, tipo_nombre = buscar_desempeno(
                        persona["nombre"], registros, nombres_equipo_completo, persona["mesa"])
                    mostrar_indicadores(registros_persona, tipo_nombre, persona["id"])
            return

        # La selección se acepta únicamente si el colaborador pertenece al alcance ya autorizado.
        seleccionado_id = str(st.query_params.get("colaborador", "") or "")
        if seleccionado_id:
            persona_seleccionada = next((c for c in equipo
                if c.get("id") == seleccionado_id and c.get("rol") != "JEFE DE CRÉDITO"), None)
            if persona_seleccionada:
                if st.button("← Volver al equipo", type="secondary"):
                    st.query_params.pop("colaborador", None)
                    st.rerun()
                st.subheader("Ficha individual de desempeño")
                mostrar_perfil_dashboard(persona_seleccionada, demo, cred_json)
                if fallo_desempeno:
                    st.warning("El desempeño no está disponible en este momento. Revisa la conexión con PRODUCTIVIDAD.")
                else:
                    registros_persona, tipo_nombre = buscar_desempeno(
                        persona_seleccionada["nombre"], registros, nombres_equipo_completo,
                        persona_seleccionada["mesa"])
                    mostrar_indicadores(registros_persona, tipo_nombre, persona_seleccionada["id"])
                return
            # Evita mantener en la URL una ficha fuera del alcance del usuario.
            st.query_params.pop("colaborador", None)
            st.rerun()
            return

        st.subheader("Dashboard del equipo")
        st.caption("Selecciona una foto para abrir la ficha del colaborador y ver sus indicadores. No se muestran datos individuales hasta entrar en la ficha.")
        if tipo_acceso == "total":
            st.caption("Vista de acceso total autorizada por LOGIN.")
        nombres_fotos = [normalizar(f["colaborador"]) for f in fotos]
        if len(nombres_fotos) != len(set(nombres_fotos)):
            st.warning("Hay fotos con el mismo nombre. No se asignan a la estructura hasta resolver el duplicado.")

        st.subheader("Filtros")
        columnas_filtros = st.columns(3)
        mesas = ["Todas"] + opciones_unicas((c["mesa"] for c in equipo), clave=clave_mesa)
        if any(not c["mesa"] for c in equipo):
            mesas.append("Sin mesa")
        with columnas_filtros[0]:
            mesa = filtro("Mesa (equipo)", mesas, "filtro_mesa")
        por_mesa = filtrar_equipo(equipo, mesa=mesa)
        jefes = ["Todos"] + opciones_unicas(c["jefe"] for c in por_mesa)
        if any(not c["jefe"] for c in por_mesa):
            jefes.append("Sin jefe asignado")
        with columnas_filtros[1]:
            jefe = filtro("Jefe de crédito", jefes, "filtro_jefe")
        por_jefe = filtrar_equipo(por_mesa, jefe=jefe)
        integrantes = [c for c in por_jefe if c["rol"] != "JEFE DE CRÉDITO"]
        roles = ["Todos"] + opciones_unicas(c["rol"] for c in integrantes)
        with columnas_filtros[2]:
            rol = filtro("Rol", roles, "filtro_rol")
        texto = st.text_input("Buscar colaborador", placeholder="Escribe el nombre de una persona")
        visibles = filtrar_equipo(integrantes, rol=rol, busqueda=texto)

        resumen = st.columns(4)
        resumen_data = [
            ("Integrantes", len({normalizar(c["nombre"]) for c in visibles})),
            ("Asesores", len({normalizar(c["nombre"]) for c in visibles if normalizar(c["rol"]) == "asesor"})),
            ("Promotores", len({normalizar(c["nombre"]) for c in visibles if normalizar(c["rol"]) == "promotor"})),
            ("Con foto", sum(1 for c in visibles if c.get("foto"))),
        ]
        for columna, (etiqueta, valor) in zip(resumen, resumen_data):
            with columna:
                st.metric(etiqueta, valor)
        if not visibles:
            st.info("No hay colaboradores con estos filtros. Ajusta la mesa, el jefe, el rol o la búsqueda.")
            return

        paginas = max(1, (len(visibles) + 11) // 12)
        pagina = filtro("Página de colaboradores", list(range(1, paginas + 1)), "pagina_dashboard")
        st.caption(f"Mostrando {(pagina - 1) * 12 + 1}–{min(pagina * 12, len(visibles))} de {len(visibles)} colaboradores.")
        visibles_pagina = visibles[(pagina - 1) * 12:pagina * 12]
        grupos_visibles = agrupar_equipos(visibles_pagina)
        for grupo in grupos_visibles:
            mesa_grupo = grupo["mesa"] or "Sin mesa asignada"
            responsable = next((c for c in equipo
                if c.get("rol") == "JEFE DE CRÉDITO"
                and clave_mesa(c.get("mesa", "")) == clave_mesa(grupo["mesa"])
                and normalizar(c.get("jefe", "")) == normalizar(grupo["jefe"])), None)
            jefe_nombre = grupo.get("jefe") or (responsable.get("nombre") if responsable else "Sin jefe asignado")
            st.markdown(
                '<div class="team-banner"><strong>' + html.escape(str(mesa_grupo)) +
                '</strong><br><span>Jefe de crédito: ' + html.escape(str(jefe_nombre)) +
                ' · ' + str(len(grupo["integrantes"])) + ' integrante(s) en esta vista</span></div>',
                unsafe_allow_html=True
            )
            columnas = st.columns(4)
            for i, colaborador in enumerate(grupo["integrantes"]):
                with columnas[i % 4]:
                    mostrar_tarjeta_colaborador(colaborador, demo, cred_json)
        st.caption("La estructura define las mesas, jefaturas y roles. Las fichas se abren desde la foto y solo se muestran si están dentro del alcance autorizado de la cuenta.")

    galeria()


if __name__ == "__main__":
    main()
